#!/usr/bin/env python3
"""Turns the personal diary workbook (NhatKy_FitJourney.xlsx) into a Supabase seed for ONE account.

    SEED_EMAIL=you@example.com python3 scripts/diary-import/build_seed.py Input/NhatKy_FitJourney.xlsx Input/generated

Writes diary-import.json (open it in the app: Profile > Import diary; the app saves it on the device and syncs it to
the cloud), plus seed.json / seed.sql (the same data in database shape, an alternative to run in Supabase > SQL Editor).
The output holds personal health data: keep it out of git (it is listed in .git/info/exclude).

Mapping (see the app's own types in src/types/index.ts):
  Chi_tiet_bua_an -> nutrition_logs: one ExtraFoodItem per eaten meal row (kcal/protein/purine/sugar; sodium mg -> salt g
                     = Na x 2.5), the three plan meal slots are marked removed because the plan menu was not followed.
  Can_nang        -> weight_entries (time of day goes into the note).
  Nhat_ky_tap     -> workout_entries.
  Ly_thuyet (D)   -> learnNotes (the sheet's personal "Ap dung cho ban" notes, one per topic; the general text is src/learn).
  Thoi_quen (A)   -> habits (the app file only; the database copy is made by the app on first sync).
"""
import json
import sys
import os
from collections import OrderedDict
from datetime import datetime, timedelta, timezone

import openpyxl

# The workbook has no daily water figures. By the owner's decision every day from the start of the diary to this date counts as
# "3 litres drunk" (the habit "Uống đủ 3 lít nước/ngày"), unless the app already has a water figure for that day.
WATER_ML = 3000
WATER_UNTIL = '2026-10-25'

from build_learn import read_topics

# The account the seed SQL is for: set SEED_EMAIL when generating (no real address is kept in the repo).
EMAIL = os.environ.get('SEED_EMAIL', 'user@example.com')
VN = timezone(timedelta(hours=7))

# Sheet meal name -> (app mealTime, typical clock time used only for ordering/display).
MEAL_TIME = {
    'Sáng': ('breakfast', (7, 30)),
    'Phụ sáng': ('morning_snack', (10, 0)),
    'Trưa': ('lunch', (12, 0)),
    'Phụ trưa': ('afternoon_snack', (14, 0)),
    'Xế': ('late_afternoon_snack', (16, 0)),
    'Tối': ('dinner', (19, 0)),
    'Phụ tối': ('evening_snack', (21, 0)),
    'Đồ uống': ('afternoon_snack', (15, 0)),
}

# Processed / sweet / fried items the sheet's own rules (Quy_tac_nhan_xet, loai "Can thay") mark for replacement.
# Beer/spirits and bamboo shoots are "can thay" there too, but they are not junk food, so they are left unflagged.
JUNK_KEYWORDS = [
    'mì tôm', 'mì indo', 'mì trộn', 'phở ăn liền', 'xúc xích', 'jambon', 'spam', 'pate', 'ruốc', 'chả lợn', 'kẹp chả',
    'bánh mì que', 'cơm nắm', 'que cay', 'latte', 'lắc sữa', 'fuzzy', 'coca', 'cold brew', 'm&m',
]


def num(v, default=0.0):
    return float(v) if isinstance(v, (int, float)) else default


def rnd(v, nd=1):
    r = round(v, nd)
    return int(r) if r == int(r) else r


def sql_str(s):
    return "'" + str(s).replace("'", "''") + "'"


def day(v):
    return v.strftime('%Y-%m-%d')


def logged_at(date, slot, index):
    _, (h, m) = MEAL_TIME[slot]
    t = datetime(date.year, date.month, date.day, h, m, tzinfo=VN) + timedelta(minutes=index)
    return t.astimezone(timezone.utc).strftime('%Y-%m-%dT%H:%M:%S.000Z')


def build(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    # ---- nutrition ----
    logs = OrderedDict()
    skipped = []
    for r in wb['Chi_tiet_bua_an'].iter_rows(min_row=5, values_only=True):
        if r[0] is None:
            continue
        date, slot, name = r[0], r[2], str(r[4] or '').strip()
        if slot not in MEAL_TIME:
            raise SystemExit(f'Unknown meal slot {slot!r} on {date}')
        kcal = num(r[5])
        if kcal == 0 and name.lower().startswith(('không ăn', 'dậy muộn')):
            # A skipped meal stays on record as a 0 kcal item so the day still counts as fully logged.
            skipped.append((day(date), slot, name))
        key = day(date)
        log = logs.setdefault(key, {'dateISO': key, 'items': []})
        n = len(log['items'])
        meal_time, _ = MEAL_TIME[slot]
        item = {
            'id': f"diary-{key.replace('-', '')}-{n + 1}",
            'name': name,
            'calories': rnd(kcal),
            'proteinG': rnd(num(r[6])),
            'loggedAtISO': logged_at(date, slot, n),
            'mealTime': meal_time,
            'purineMg': rnd(num(r[7])),
            'saltG': rnd(num(r[9]) * 2.5 / 1000, 4),
            'sugarG': rnd(num(r[8])),
        }
        if any(k in name.lower() for k in JUNK_KEYWORDS):
            item['junk'] = True
        log['items'].append(item)

    nutrition = []
    for key, log in logs.items():
        items = log['items']
        nutrition.append({
            'entry_date': key,
            'meals_eaten': {},
            'extra_calories': rnd(sum(i['calories'] for i in items)),
            'extra_protein_g': rnd(sum(i['proteinG'] for i in items)),
            'extra_items': items,
            'water_ml': 0,
            'removed_meal_slots': ['meal1', 'meal2', 'meal3'],
        })

    # ---- weight ----
    weights = []
    for r in wb['Can_nang'].iter_rows(min_row=5, values_only=True):
        if not isinstance(r[1], (int, float)) or r[0] is None:
            continue
        when, note = r[2], (r[4] or '').strip()
        bits = []
        if when and when != 'Không rõ' and when.lower() not in note.lower():
            bits.append(f'Cân {when.lower()}')
        if note:
            bits.append(note)
        weights.append({'entry_date': day(r[0]), 'weight_kg': rnd(r[1], 2), 'note': ' · '.join(bits) or None})
        
    # ---- workouts ----
    workouts = []
    for r in wb['Nhat_ky_tap'].iter_rows(min_row=5, values_only=True):
        if r[0] is None or not r[1]:
            continue
        unit = str(r[3] or '')
        if unit != 'phút':
            raise SystemExit(f'Workout unit {unit!r} on {day(r[0])} is not minutes; extend the mapper')
        workouts.append({
            'client_id': f"diary-{day(r[0]).replace('-', '')}-1",
            'entry_date': day(r[0]),
            'template_id': None,
            'label': str(r[1]),
            'duration_min': rnd(num(r[2]), 1),
            'calories_burned': rnd(num(r[6])),
        })

    # ---- profile (what the workbook states; unknowns stay NULL = left untouched) ----
    thong_so = {str(row[0]): row[1] for row in wb['Thong_so'].iter_rows(min_row=5, values_only=True) if row[0]}
    tdee = num(thong_so.get('TDEE cơ sở hiện tại'))
    deficit = num(thong_so.get('Thâm hụt TB – mức ĐẠT (≥)'))
    protein = num(thong_so.get('Protein – mức ĐẠT (≥)'))
    first_weight = weights[0]['weight_kg']
    profile = {
        'start_weight_kg': first_weight,
        'start_date': weights[0]['entry_date'],
        'calorie_target_kcal': rnd(tdee - deficit),
        'target_daily_deficit_kcal': rnd(deficit),
        'protein_target_g': rnd(protein),
    }
    return {'email': EMAIL, 'profile': profile, 'weights': weights, 'nutrition': nutrition, 'workouts': workouts,
            'skipped_meals': skipped, 'habits': read_habits(wb), 'learn_notes': read_topics(wb)[1], 'uric': read_uric(wb)}


# Thoi_quen section A: habit name keyword -> (tracked day by day in section B, auto-measured from app data)
HABIT_RULES = [
    ('không húp nước dùng', True, None),
    ('gà bỏ da', True, None),
    ('ước tính kcal', True, None),
    ('rau trước', True, None),
    ('ít ăn ngọt', True, None),
    ('giấm táo', False, None),
    ('3 lít nước', True, 'water'),
    ('snack ăn kiêng', True, None),
    ('cân hàng ngày', True, 'weigh'),
    ('đọc nhãn', False, None),
    ('đi bộ sau bữa', True, 'walk'),
]
HABIT_STATUS = {'Đã thành nếp': 'formed', 'Đang xây': 'building', 'Cần xây': 'todo'}


def read_habits(wb):
    habits = []
    for r in wb['Thoi_quen'].iter_rows(min_row=5, max_row=30, values_only=True):
        if not r[0] or r[1] not in HABIT_STATUS:
            continue
        name = str(r[0]).strip()
        rule = next((x for x in HABIT_RULES if x[0] in name.lower()), None)
        if rule is None:
            raise SystemExit(f'Habit {name!r} has no tracking rule in HABIT_RULES')
        habits.append({
            'id': f'diary-h-{len(habits) + 1:02d}',
            'name': name,
            'status': HABIT_STATUS[r[1]],
            'assessment': str(r[2] or '').strip(),
            'kaizen': str(r[3] or '').strip(),
            'measuredBy': '' if str(r[4] or '').strip() in ('', '—') else str(r[4]).strip(),
            'tracked': rule[1],
            'auto': rule[2],
            'sort': len(habits) + 1,
            'archived': False,
        })
    assert len(habits) == len(HABIT_RULES), f'expected {len(HABIT_RULES)} habits, found {len(habits)}'
    assert sum(h['tracked'] for h in habits) == 9, 'section B tracks 9 habits'
    return habits


# The workbook records the uric acid test only inside its notes (Tong_hop_ngay 16/10: "Uric 457 µmol/L; steak tối"; Phan_tich;
# Thong_so). There is no table of tests, so the one known reading is stated here and checked against the workbook text.
URIC_READINGS = [{'dateKey': '2026-10-16', 'valueUmol': 457, 'note': 'Kết quả xét nghiệm ghi trong nhật ký, đã hạ so với lần trước.'}]


def read_uric(wb):
    texts = [str(c) for ws in wb for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    for i, reading in enumerate(URIC_READINGS):
        assert any(f"{reading['valueUmol']} µmol/L" in t for t in texts), f"uric reading {reading['valueUmol']} not found in the workbook"
        reading['id'] = f'diary-u-{i + 1}'
    return URIC_READINGS


def noon(date):  # noon in Vietnam = 05:00Z, the same calendar day in UTC and in UTC+7
    return f'{date}T05:00:00.000Z'


def weight_ids(d):
    """App id and time of each weigh-in: diary-w-YYYYMMDD-n, n counting that day's readings."""
    out = []
    for i, w in enumerate(d['weights']):
        n = sum(1 for x in d['weights'][:i] if x['entry_date'] == w['entry_date']) + 1
        out.append((f"diary-w-{w['entry_date'].replace('-', '')}-{n}", noon(w['entry_date'])))
    return out


def water_days(d):
    """Every calendar day from the first diary day to WATER_UNTIL, as YYYY-MM-DD."""
    start = datetime.strptime(d['profile']['start_date'], '%Y-%m-%d')
    end = datetime.strptime(WATER_UNTIL, '%Y-%m-%d')
    return [(start + timedelta(days=i)).strftime('%Y-%m-%d') for i in range((end - start).days + 1)]


def to_app_file(d):
    """The same data in the shape the app imports (src/import/diaryImport.ts)."""

    p = d['profile']
    return {
        'format': 'fitjourney-diary',
        'version': 1,
        'profile': {
            'startWeightKg': p['start_weight_kg'],
            'startDateISO': noon(p['start_date']),
            'calorieTargetKcal': p['calorie_target_kcal'],
            'targetDailyDeficitKcal': p['target_daily_deficit_kcal'],
            'proteinTargetG': p['protein_target_g'],
        },
        'weights': [
            {'id': wid, 'dateISO': iso, 'weightKg': w['weight_kg'], **({'note': w['note']} if w['note'] else {})}
            for w, (wid, iso) in zip(d['weights'], weight_ids(d))
        ],
        'nutrition': [{'dateISO': n['entry_date'], 'items': n['extra_items']} for n in d['nutrition']],
        'workouts': [
            {'id': w['client_id'], 'dateISO': noon(w['entry_date']), 'templateId': None, 'label': w['label'],
             'durationMin': w['duration_min'], 'caloriesBurned': w['calories_burned']}
            for w in d['workouts']
        ],
        'habits': d['habits'],
        'learnNotes': d['learn_notes'],
        'uric': d['uric'],
        'water': {'fromDateKey': water_days(d)[0], 'toDateKey': WATER_UNTIL, 'ml': WATER_ML},
    }


def to_sql(d):
    p = d['profile']
    out = [
        f"-- Seed nhật ký FitJourney cho {d['email']} (tạo bởi scripts/diary-import/build_seed.py).",
        '-- Chạy trong Supabase > SQL Editor. Tài khoản phải đã đăng ký (có trong auth.users). Chạy lại được: ghi đè đúng các ngày này.',
        '-- Chỗ NULL ở khối "Hồ sơ" = giữ nguyên giá trị đang có. Điền nếu bạn muốn đặt.',
        'do $seed$',
        'declare',
        '  uid uuid;',
        '  -- Hồ sơ: điền các giá trị bạn muốn đặt (để NULL = không đổi).',
        '  v_goal_weight numeric := NULL;      -- cân nặng mục tiêu (kg)',
        '  v_height numeric := NULL;           -- chiều cao (cm)',
        '  v_age integer := NULL;              -- tuổi',
        "  v_gender text := NULL;              -- 'male' | 'female'",
        "  v_activity text := NULL;            -- 'sedentary' | 'light' | 'moderate' | 'active' | 'very_active'",
        '  v_water_ml numeric := NULL;         -- mục tiêu nước/ngày (ml)',
        'begin',
        f"  select id into uid from auth.users where lower(email) = lower({sql_str(d['email'])});",
        f"  if uid is null then raise exception 'Chưa có tài khoản %. Hãy đăng ký/đăng nhập app bằng email này trước.', {sql_str(d['email'])}; end if;",
        '  insert into public.profiles (id, email) values (uid, ' + sql_str(d['email']) + ') on conflict (id) do nothing;',
        '  update public.profiles set',
        f"    start_weight_kg = {p['start_weight_kg']}, start_date = {sql_str(p['start_date'])},",
        f"    calorie_target_kcal = {p['calorie_target_kcal']}, target_daily_deficit_kcal = {p['target_daily_deficit_kcal']},",
        f"    protein_target_g = {p['protein_target_g']},",
        '    goal_weight_kg = coalesce(v_goal_weight, goal_weight_kg), height_cm = coalesce(v_height, height_cm),',
        '    age = coalesce(v_age, age), gender = coalesce(v_gender, gender),',
        '    activity_level = coalesce(v_activity, activity_level), water_target_ml = coalesce(v_water_ml, water_target_ml),',
        '    updated_at = now()',
        '  where id = uid;',
        '',
        '  -- Cân nặng (khóa client_id, giống app: chạy lại hoặc nhập thêm trong app không nhân đôi). Cần đã chạy schema.sql bản mới (cột client_id, logged_at).',
    ]
    for w, (wid, iso) in zip(d['weights'], weight_ids(d)):
        note = sql_str(w['note']) + '::text' if w['note'] else 'NULL::text'
        day_lit = sql_str(w['entry_date'])
        # A row an earlier seed (or app version) saved without an id is adopted, so nothing is stored twice.
        out.append(
            f"  update public.weight_entries set client_id = {sql_str(wid)}, logged_at = {sql_str(iso)}::timestamptz where id = (select id from public.weight_entries "
            f"where user_id = uid and client_id is null and entry_date = {day_lit} and weight_kg = {w['weight_kg']} limit 1) "
            f"and not exists (select 1 from public.weight_entries where user_id = uid and client_id = {sql_str(wid)});"
        )
        out.append(
            f"  insert into public.weight_entries (user_id, client_id, entry_date, logged_at, weight_kg, note) values "
            f"(uid, {sql_str(wid)}, date {day_lit}, {sql_str(iso)}::timestamptz, {w['weight_kg']}::numeric, {note}) "
            f"on conflict (user_id, client_id) do update set entry_date = excluded.entry_date, logged_at = excluded.logged_at, weight_kg = excluded.weight_kg, note = excluded.note;"
        )
        out.append(f"  delete from public.weight_entries where user_id = uid and client_id is null and entry_date = {day_lit} and weight_kg = {w['weight_kg']};")
    out += ['', '  -- Buổi tập (khóa client_id nên chạy lại không nhân đôi).']
    for w in d['workouts']:
        out.append(
            f"  insert into public.workout_entries (user_id, client_id, entry_date, template_id, label, duration_min, calories_burned) values "
            f"(uid, {sql_str(w['client_id'])}, {sql_str(w['entry_date'])}, NULL, {sql_str(w['label'])}, {w['duration_min']}, {w['calories_burned']}) "
            f"on conflict (user_id, client_id) do update set entry_date = excluded.entry_date, label = excluded.label, "
            f"duration_min = excluded.duration_min, calories_burned = excluded.calories_burned;"
        )
    out += ['', '  -- Nhật ký ăn uống theo ngày (ghi đè các ngày này).']
    for n in d['nutrition']:
        items = json.dumps(n['extra_items'], ensure_ascii=False)
        out.append(
            f"  insert into public.nutrition_logs (user_id, entry_date, meals_eaten, extra_calories, extra_protein_g, extra_items, water_ml, removed_meal_slots) values "
            f"(uid, {sql_str(n['entry_date'])}, '{{}}'::jsonb, {n['extra_calories']}, {n['extra_protein_g']}, {sql_str(items)}::jsonb, 0, "
            f"{sql_str(json.dumps(n['removed_meal_slots']))}::jsonb) "
            f"on conflict (user_id, entry_date) do update set meals_eaten = excluded.meals_eaten, extra_calories = excluded.extra_calories, "
            f"extra_protein_g = excluded.extra_protein_g, extra_items = excluded.extra_items, removed_meal_slots = excluded.removed_meal_slots, updated_at = now();"
        )
    out += ['', f'  -- Nước: các ngày từ đầu nhật ký đến {WATER_UNTIL} tính là đủ {WATER_ML} ml, trừ ngày đã có số nước.']
    for day_key in water_days(d):
        out.append(
            f"  insert into public.nutrition_logs (user_id, entry_date, water_ml) values (uid, {sql_str(day_key)}, {WATER_ML}) "
            f"on conflict (user_id, entry_date) do update set water_ml = {WATER_ML}, updated_at = now() where public.nutrition_logs.water_ml = 0;"
        )
    out += ['end', '$seed$;', '']
    return '\n'.join(out)


if __name__ == '__main__':
    src, dest = sys.argv[1], sys.argv[2]
    data = build(src)
    # Self-checks so a wrong sheet layout fails loudly instead of seeding rubbish.
    assert data['nutrition'], 'no nutrition rows'
    for n in data['nutrition']:
        assert abs(sum(i['calories'] for i in n['extra_items']) - n['extra_calories']) < 0.01
        assert len({i['id'] for i in n['extra_items']}) == len(n['extra_items'])
    with open(f'{dest}/seed.json', 'w') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    with open(f'{dest}/diary-import.json', 'w') as f:
        json.dump(to_app_file(data), f, ensure_ascii=False, indent=1)
    with open(f'{dest}/seed.sql', 'w') as f:
        f.write(to_sql(data))
    kcal = sum(n['extra_calories'] for n in data['nutrition'])
    print(f"days={len(data['nutrition'])} items={sum(len(n['extra_items']) for n in data['nutrition'])} kcal={kcal} "
          f"weights={len(data['weights'])} workouts={len(data['workouts'])} skipped={data['skipped_meals']}")
